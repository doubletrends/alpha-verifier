"""Stage 3 command: select the strongest skew sheets from Stage 2 artifacts."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

from engine import selection, shift
from universe import all_in_family, all_nodes, load_universe
from workspace import BASELINE_NODE, Workspace

ROOT = Path(__file__).resolve().parents[1]


def _copy_selected_workbook_sheet(source: Path, target: Path, bin_index: int) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    wb = load_workbook(target)
    keep = wb.sheetnames[int(bin_index)]
    for sheet_name in list(wb.sheetnames):
        if sheet_name != keep:
            del wb[sheet_name]
    wb.save(target)
    wb.close()


def cmd_selection(ws: Workspace, family: str | None = None) -> None:
    """Stage 3: score every node/bin shift sheet and select the global top-k."""
    universe = load_universe(ws.universe_path)
    nodes = [
        n for n in (all_in_family(universe, family) if family else all_nodes(universe))
        if n["id"] != BASELINE_NODE and ws.has_shift_cube(n["family"], n["id"])
    ]
    if not nodes:
        print("No shift arrays to select from - run --shift first.")
        return

    def load_cube(node: dict) -> dict:
        return shift.load(ws.shift_cube_path(node["family"], node["id"]))

    result = selection.rank_shift_sheets(
        nodes,
        load_cube,
        ws.min_dev,
        ws.min_bin_n,
        ws.min_run,
        ws.selection_top_k,
    )
    result["workspace"] = ws.dir.name
    result["families"] = sorted({n["family"] for n in nodes})

    selected = result["selected"]
    copied_npz = 0
    copied_xlsx = 0
    by_id = {n["id"]: n for n in nodes}
    for row in selected:
        node = by_id[row["node"]]
        source = load_cube(node)
        sheet = selection.sheet_from_shift_cube(source, row)
        meta = {
            **sheet["meta"],
            "artifact": "03_selection_array",
            "source": str(ws.shift_cube_path(row["family"], row["node"]).relative_to(ROOT)),
            "node": row["node"],
            "family": row["family"],
            "feature": row["feature"],
            "params": row["params"],
            "selection": {
                "rank": row["rank"],
                "score": row["score"],
                "method": result["method"]["score"],
            },
            "economic": row["economic"],
        }
        selection.save_sheet(sheet, ws.selection_array_path(row), meta)
        copied_npz += 1

        source_xlsx = ws.shift_surface_path(node["family"], node["id"])
        target_xlsx = ws.selection_surface_path(row)
        try:
            _copy_selected_workbook_sheet(source_xlsx, target_xlsx, int(row["bin"]))
            copied_xlsx += 1
        except PermissionError:
            print(f"  rank {row['rank']:>3} {row['node']:<26} [locked] close it in Excel and re-run")

    ws.write_json(ws.selection_path, result)

    print(f"\n=== 3. Selection [{ws.dir.name}] ===")
    print(
        f"  scored {len(result['candidates'])} node/bin sheets from 02_shift_array; "
        f"selected top {len(selected)} by upside-vs-downside skew"
    )
    print(f"  score = max |shift(+Δ,t) - shift(-Δ,t)|, min bin n = {ws.min_bin_n}\n")
    if selected:
        print(f"  {'rank':>4}  {'node':<26}{'family':<13}{'bin':>5}  {'score':>7}  {'econ':>5}  best skew cell")
        print(f"  {'-'*4}  {'-'*26}{'-'*13}{'-'*5}  {'-'*7}  {'-'*5}  {'-'*42}")
        for row in selected:
            c = row["best_cell"]
            econ = "yes" if row["economic"]["passed"] else "no"
            print(
                f"  {row['rank']:>4}  {row['node']:<26}{row['family']:<13}"
                f"{row['bin_number']:>5}  {row['score']:>6.1f}  {econ:>5}  "
                f"|Δ|={c['Δ_abs']:+.0%} +{c['horizon']}{ws.horizon_unit} "
                f"skew={c['skew']:+.1f}pp n={c['bin_n']}"
            )

    print(f"\n  wrote {copied_npz} .npz artifacts under {ws.dir.relative_to(ROOT)}/03_selection_array/")
    print(f"  wrote {copied_xlsx} workbooks under {ws.dir.relative_to(ROOT)}/03_selection_xlsx/")
