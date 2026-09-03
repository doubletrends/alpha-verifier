"""
The pipeline's deliverable.

write_barrier_xlsx renders a cube as one tab per condition bin. Each tab is that bin's
whole (θ x horizon) face, so the workbook holds every value the cube holds -- it is
a faithful view of the measurement, not a summary of it.

Cells are the raw conditional probability. Nothing is subtracted, so a cell reads on its
own terms -- "under this condition a -5% touch within 7 days happens 50% of the time" --
and the colour bands show where the probability actually lives rather than where it
differs from an average the reader cannot see.
"""

from __future__ import annotations

import re
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine.barrier import MIN_BIN_N

_PCT_FMT = '0.0%'
_PP_FMT = '+0.0;-0.0;0.0'
_PVAL_FMT = '0.0000'

_WHITE, _AMBER, _RED = 'FFFFFF', 'FFD166', 'C00000'
_GREEN = '00875A'
_FILL_ROW1 = PatternFill('solid', start_color='666666', end_color='666666')
_FILL_ROW2 = PatternFill('solid', start_color='B2B2B2', end_color='B2B2B2')
_FILL_ROW3 = PatternFill('solid', start_color='CCCCCC', end_color='CCCCCC')
_FILL_NBAND = PatternFill('solid', start_color='EFEFEF', end_color='EFEFEF')
_FILL_NA = PatternFill('solid', start_color='F7F7F7', end_color='F7F7F7')
_FILL_MID = PatternFill('solid', start_color='DDDDDD', end_color='DDDDDD')
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

_COL_ROW  = 5   # horizon labels
_N_ROW    = 6   # observations behind this bin at each horizon
_DATA_ROW = 7   # first θ row

_ABBREV = {'rsi', 'ma', 'atr', 'dxy', 'macd', 'bb', 'mvrv', 'wr', 'roc', 'vol'}


def feature_label(node_id: str) -> str:
    words = []
    for p in node_id.split('_'):
        if p.isdigit():
            if words:
                words[-1] += '-' + p
        elif p.lower() in _ABBREV:
            words.append(p.upper())
        else:
            words.append(p.title())
    return ' '.join(words)


def _sheet_names(labels: list[str]) -> list[str]:
    """
    Tab names are the conditions themselves -- 'x < 0.12', '0.12 < x < 0.34' -- so the
    tab strip reads as the ladder of bins with nothing to decode.

    Excel caps a sheet name at 31 chars and forbids : \\ / ? * [ ]. Truncation could in
    principle collide two long labels, so a numeric suffix is appended only when that
    actually happens rather than pre-emptively numbering every tab.
    """
    out: list[str] = []
    for label in labels:
        name = re.sub(r'[:\\/?*\[\]]', '-', label)[:31]
        if name in out:
            base = name[:28]
            k = 2
            while f'{base}~{k}' in out:
                k += 1
            name = f'{base}~{k}'
        out.append(name)
    return out


def _write_headers(ws, row1: str, row2: str, row3: str, merge_end: str) -> None:
    for r, text, font, fill in [
        (1, row1, Font(bold=True, size=14, color='FFFFFF'), _FILL_ROW1),
        (2, row2, Font(bold=True, size=11, color='000000'), _FILL_ROW2),
        (3, row3, Font(size=10, color='000000'), _FILL_ROW3),
    ]:
        ws.merge_cells(f'A{r}:{merge_end}{r}')
        c = ws[f'A{r}']
        c.value, c.font, c.fill, c.alignment = text, font, fill, _CENTER
        ws.row_dimensions[r].height = 18


def write_barrier_xlsx(
    cube:    dict,
    path:    Path,
    node_id: str,
    feature: str,
    params:  dict,
    unit:    str = 'd',
) -> None:
    """
    One tab per condition bin; each tab is that bin's full θ x horizon face.

    Rows run from the highest barrier at the top to the lowest at the bottom, the way a
    price ladder reads: up the sheet is up in price. θ = 0 sits in the middle and is
    shaded, marking the boundary between two different questions -- above it a cell asks
    whether the *high* reached that level, below it whether the *low* did.

    The colour scale is a fixed 0..100% on every tab, so flipping between them shows the
    probability band moving with the condition instead of each tab being rescaled to
    look alike.

    The n band under the header is the observation count behind that bin at each
    horizon; it falls as h grows, because the last h bars have no realized forward
    window.
    """
    prob     = cube['prob']
    thetas   = cube['thetas']
    horizons = cube['horizons']
    bin_n    = cube['bin_n']
    labels   = cube['meta']['bin_labels']
    n_th, n_bins, n_h = prob.shape

    order = np.argsort(thetas)[::-1]        # highest barrier on the top row
    end   = get_column_letter(1 + n_h)

    wb = Workbook()
    wb.remove(wb.active)

    unconditional = n_bins == 1 and labels[0] == 'all'
    names = _sheet_names(list(labels))

    for b in range(n_bins):
        ws = wb.create_sheet(names[b])
        title = (f'{feature_label(node_id)} — unconditional, every bar'
                 if unconditional else
                 f'{feature_label(node_id)} — condition: {feature} in {labels[b]}')
        _write_headers(
            ws,
            title,
            ('P(price touches θ within h)' if unconditional
             else 'P(price touches θ within h | condition)'),
            f'node {node_id}   params={params}   bin {b + 1} of {n_bins}   '
            f'rows: barrier θ, highest at the top   columns: horizon   '
            f'intraday high/low, window opens at t+1',
            merge_end=end,
        )

        c = ws.cell(row=_COL_ROW, column=1, value='θ')
        c.font, c.alignment = Font(bold=True), _CENTER
        for j, h in enumerate(horizons):
            c = ws.cell(row=_COL_ROW, column=2 + j, value=f'+{int(h)}{unit}')
            c.font, c.alignment = Font(bold=True), _CENTER

        c = ws.cell(row=_N_ROW, column=1, value='n =')
        c.font, c.fill = Font(bold=True, italic=True, size=9), _FILL_NBAND
        for j in range(n_h):
            c = ws.cell(row=_N_ROW, column=2 + j, value=int(bin_n[b, j]))
            c.font, c.fill, c.alignment = Font(italic=True, size=9), _FILL_NBAND, _CENTER

        for r_off, i in enumerate(order):
            r = _DATA_ROW + r_off
            th = float(thetas[i])
            is_zero = abs(th) < 1e-12
            tc = ws.cell(row=r, column=1, value=th)
            tc.number_format = '+0%;-0%;0%'
            tc.font, tc.alignment = Font(bold=True), _CENTER
            if is_zero:
                tc.fill = _FILL_MID
            for j in range(n_h):
                v = prob[i, b, j]
                cell = ws.cell(row=r, column=2 + j,
                               value=None if not np.isfinite(v) else round(float(v), 4))
                cell.number_format = _PCT_FMT
                if not np.isfinite(v):
                    cell.fill = _FILL_NA

        ws.conditional_formatting.add(
            f'B{_DATA_ROW}:{end}{_DATA_ROW + n_th - 1}',
            ColorScaleRule(start_type='num', start_value=0,   start_color=_WHITE,
                           mid_type='num',   mid_value=0.5,   mid_color=_AMBER,
                           end_type='num',   end_value=1.0,   end_color=_RED))

        ws.column_dimensions['A'].width = 8
        for j in range(n_h):
            ws.column_dimensions[get_column_letter(2 + j)].width = 6.5
        ws.freeze_panes = f'B{_DATA_ROW}'

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_skew_xlsx(
    result:  dict,
    path:    Path,
    node_id: str,
    feature: str,
    params:  dict,
    unit:    str = 'd',
) -> None:
    """
    Human-readable skew artifact: the mirrored grid's up/down asymmetry.

    The first sheet is the per-horizon peak excess skew -- the search statistic --validation
    later null-tests. The remaining sheets render, per condition bin, the conditional
    skew (P(+theta) - P(-theta), still carrying the asset's drift) and the excess skew
    (that minus the baseline node's own skew, so drift is removed).

    Positive means the condition reaches *up* more readily than down. Rows are barrier
    magnitude |theta|, largest at the top.
    """
    cond   = result['cond_skew']
    excess = result['excess_skew']
    mags   = result['mags']
    horizons = result['horizons']
    bin_n  = result['bin_n']
    labels = result['meta']['bin_labels']
    n_mag, n_bins, n_h = excess.shape

    wb = Workbook()
    ws = wb.active
    ws.title = 'summary'
    end = 'F'
    _write_headers(
        ws,
        f'{feature_label(node_id)} — skew summary',
        'Peak excess skew by horizon; this is the statistic --validation null-tests',
        f'node {node_id}   params={params}   '
        f'excess skew = P(+θ|bin) − P(−θ|bin) minus the baseline node’s same difference',
        merge_end=end,
    )
    headers = ['horizon', 'peak excess pp', 'at |θ|', 'bin', 'cond skew pp', 'n']
    for col, header in enumerate(headers, 1):
        c = ws.cell(row=_COL_ROW, column=col, value=header)
        c.font, c.alignment = Font(bold=True), _CENTER
    for j, h in enumerate(horizons):
        face = np.abs(excess[:, :, j]).copy()
        for b in range(n_bins):
            if bin_n[b, j] < MIN_BIN_N:
                face[:, b] = np.nan
        r = _COL_ROW + 1 + j
        if np.isfinite(face).any():
            k = int(np.nanargmax(face))
            mi, bi = divmod(k, n_bins)
            row = [f'+{int(h)}{unit}', float(excess[mi, bi, j]), float(mags[mi]),
                   bi + 1, float(cond[mi, bi, j]), int(bin_n[bi, j])]
        else:
            row = [f'+{int(h)}{unit}', None, None, None, None, None]
        for col, v in enumerate(row, 1):
            c = ws.cell(row=r, column=col, value=v)
            if col in (2, 5):
                c.number_format = _PP_FMT
            if col == 3:
                c.number_format = '0.0%'
            c.alignment = _CENTER
    for col, width in enumerate([10, 15, 9, 8, 13, 8], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f'A{_COL_ROW + 1}'

    used = {'summary'}
    names = _sheet_names(list(labels))
    for b in range(n_bins):
        common = (f'node {node_id}   bin {b + 1} of {n_bins}: {labels[b]}   '
                  'rows: barrier magnitude |θ|, largest at the top   columns: horizon')
        _write_validation_matrix(
            wb, used, f'skew {names[b]}',
            cond[:, b, :], mags, horizons,
            f'{feature_label(node_id)} — conditional skew',
            'P(+θ | condition) − P(−θ | condition), percentage points; carries drift',
            common,
            _PP_FMT,
            unit,
            pvalue_scale=False,
        )
        _write_validation_matrix(
            wb, used, f'excess {names[b]}',
            excess[:, b, :], mags, horizons,
            f'{feature_label(node_id)} — excess skew',
            'Conditional skew minus the baseline node’s skew; drift removed',
            common,
            _PP_FMT,
            unit,
            pvalue_scale=False,
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _safe_sheet_name(name: str, used: set[str]) -> str:
    name = re.sub(r'[:\\/?*\[\]]', '-', name)[:31]
    if name not in used:
        used.add(name)
        return name
    base = name[:28]
    k = 2
    while f'{base}~{k}' in used:
        k += 1
    name = f'{base}~{k}'
    used.add(name)
    return name


def _write_validation_matrix(
    wb: Workbook,
    used: set[str],
    sheet_name: str,
    values: np.ndarray,
    thetas: np.ndarray,
    horizons: np.ndarray,
    title: str,
    subtitle: str,
    row3: str,
    number_format: str,
    unit: str,
    pvalue_scale: bool,
) -> None:
    n_th, n_h = values.shape
    order = np.argsort(thetas)[::-1]
    end = get_column_letter(1 + n_h)
    ws = wb.create_sheet(_safe_sheet_name(sheet_name, used))
    _write_headers(ws, title, subtitle, row3, merge_end=end)

    c = ws.cell(row=_COL_ROW, column=1, value='θ')
    c.font, c.alignment = Font(bold=True), _CENTER
    for j, h in enumerate(horizons):
        c = ws.cell(row=_COL_ROW, column=2 + j, value=f'+{int(h)}{unit}')
        c.font, c.alignment = Font(bold=True), _CENTER

    for r_off, i in enumerate(order):
        r = _DATA_ROW + r_off
        th = float(thetas[i])
        tc = ws.cell(row=r, column=1, value=th)
        tc.number_format = '+0%;-0%;0%'
        tc.font, tc.alignment = Font(bold=True), _CENTER
        if abs(th) < 1e-12:
            tc.fill = _FILL_MID
        for j in range(n_h):
            v = values[i, j]
            cell = ws.cell(row=r, column=2 + j,
                           value=None if not np.isfinite(v) else float(v))
            cell.number_format = number_format
            if not np.isfinite(v):
                cell.fill = _FILL_NA

    rng = f'B{_DATA_ROW}:{end}{_DATA_ROW + n_th - 1}'
    if pvalue_scale:
        ws.conditional_formatting.add(
            rng,
            ColorScaleRule(start_type='num', start_value=0, start_color=_GREEN,
                           mid_type='num', mid_value=0.05, mid_color=_AMBER,
                           end_type='num', end_value=1.0, end_color=_RED))
    else:
        ws.conditional_formatting.add(
            rng,
            ColorScaleRule(start_type='num', start_value=-20, start_color=_GREEN,
                           mid_type='num', mid_value=0, mid_color=_WHITE,
                           end_type='num', end_value=20, end_color=_RED))

    ws.column_dimensions['A'].width = 8
    for j in range(n_h):
        ws.column_dimensions[get_column_letter(2 + j)].width = 8.5
    ws.freeze_panes = f'B{_DATA_ROW}'


def write_validation_xlsx(
    result:  dict,
    path:    Path,
    node_id: str,
    feature: str,
    params:  dict,
    unit:    str = 'd',
) -> None:
    """
    Human-readable validation artifact.

    The first sheet is the horizon-level search statistic used by the gate. The
    remaining sheets render the per-cell signed deviation and pointwise p-value for
    each condition bin. Pointwise p-values are for reading the surface; discoveries are
    still assigned only by --gate across the whole sweep.
    """
    thetas = result['thetas']
    horizons = result['horizons']
    labels = result['meta']['bin_labels']
    n_bins = result['cell_real'].shape[1]

    wb = Workbook()
    ws = wb.active
    ws.title = 'summary'
    end = 'F'
    _write_headers(
        ws,
        f'{feature_label(node_id)} — validation summary',
        'Peak null test by horizon; this is the statistic --gate corrects',
        f'node {node_id}   params={params}   p floor = 1/(usable shifts + 1)',
        merge_end=end,
    )
    headers = ['horizon', 'peak dev pp', 'null p95 pp', 'peak p', 'usable shifts', 'p floor']
    for col, header in enumerate(headers, 1):
        c = ws.cell(row=_COL_ROW, column=col, value=header)
        c.font, c.alignment = Font(bold=True), _CENTER
    for r_off, h in enumerate(horizons):
        r = _COL_ROW + 1 + r_off
        n = int(result['n_shifts'][r_off])
        floor = None if n <= 0 else 1.0 / (1.0 + n)
        row = [
            f'+{int(h)}{unit}',
            result['peak_real'][r_off],
            result['peak_p95'][r_off],
            result['peak_p'][r_off],
            n,
            floor,
        ]
        for col, v in enumerate(row, 1):
            c = ws.cell(row=r, column=col,
                        value=None if isinstance(v, float) and not np.isfinite(v) else v)
            if col in (2, 3):
                c.number_format = _PP_FMT
            if col in (4, 6):
                c.number_format = _PVAL_FMT
            c.alignment = _CENTER
    for col, width in enumerate([10, 13, 13, 10, 14, 10], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f'A{_COL_ROW + 1}'

    used = {'summary'}
    names = _sheet_names(list(labels))
    for b in range(n_bins):
        label = labels[b]
        suffix = names[b]
        common = (f'node {node_id}   bin {b + 1} of {n_bins}: {label}   '
                  'rows: barrier theta   columns: horizon')
        _write_validation_matrix(
            wb, used, f'dev {suffix}',
            result['cell_real'][:, b, :], thetas, horizons,
            f'{feature_label(node_id)} — signed validation deviation',
            'Deviation from this node bin sample rate, percentage points',
            common,
            _PP_FMT,
            unit,
            pvalue_scale=False,
        )
        _write_validation_matrix(
            wb, used, f'p {suffix}',
            result['cell_p'][:, b, :], thetas, horizons,
            f'{feature_label(node_id)} — pointwise validation p-values',
            'Per-cell null p-values; useful for reading, not for discovery claims',
            common,
            _PVAL_FMT,
            unit,
            pvalue_scale=True,
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
