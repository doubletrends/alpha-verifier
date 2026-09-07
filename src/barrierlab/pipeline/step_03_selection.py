"""Stage 3 command: select nodes by their strongest-bin probability skew."""

from __future__ import annotations

from datetime import datetime, timezone
import shutil
from pathlib import Path

from openpyxl import load_workbook

from barrierlab.domain import selection
from barrierlab.infrastructure import artifact_io
from barrierlab.infrastructure.workspace import BASELINE_NODE, Workspace
from barrierlab.pipeline.reporting import StageReport


def _copy_selected_workbook_sheet(source: Path, target: Path, bin_index: int) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    wb = load_workbook(target)
    keep = wb.sheetnames[int(bin_index)]
    for sheet_name in list(wb.sheetnames):
        if sheet_name != keep:
            del wb[sheet_name]
    wb.save(target)
    wb.close()


def _clear_selection_artifacts(ws: Workspace) -> None:
    """Discard generated Stage 3 views before replacing the selection manifest."""
    for file_type in ("array", "spreadsheet", "plot"):
        directory = ws.selection_path.parent / file_type
        if not directory.exists():
            continue
        for path in directory.iterdir():
            if path.is_file():
                path.unlink()


def cmd_selection(ws: Workspace) -> None:
    """Stage 3: retain nodes with the largest single-bin probability skew."""
    nodes = [
        n for n in ws.catalog.all_nodes()
        if n["id"] != BASELINE_NODE and ws.has_shift_cube(n["id"])
    ]
    report = StageReport(3, "select", ws.dir.name)
    if not nodes:
        report.line("no shift arrays available; run compare first")
        report.completed()
        return

    def load_cube(node: dict) -> dict:
        return artifact_io.load_shift(ws.shift_cube_path(node["id"]))

    baseline_path = ws.shift_cube_path(BASELINE_NODE)
    if not baseline_path.exists():
        report.line("no baseline shift array available; run compare first")
        report.completed()
        return
    delta, horizon = ws.target(artifact_io.load_shift(baseline_path)["base"])
    result = {
        "generated": datetime.now(timezone.utc).isoformat(),
        "artifact": "03_selection",
        "source": "02_shift",
        **selection.rank_nodes(
            nodes,
            load_cube,
            ws.selection_top_k,
            delta,
            horizon,
        ),
    }
    result["workspace"] = ws.dir.name
    result["families"] = sorted({n["family"] for n in nodes})

    selected = result["selected"]
    report.line(
        f"scoring {len(result['candidates'])} condition bins from {len(nodes)} nodes · "
        f"retaining top {len(selected)} by absolute skew"
    )
    _clear_selection_artifacts(ws)
    copied_npz = 0
    copied_xlsx = 0
    warnings = []
    by_id = {n["id"]: n for n in nodes}
    for row in selected:
        node = by_id[row["node"]]
        source = load_cube(node)
        selected_node = selection.selected_node_from_shift_cube(source, row)
        meta = {
            **selected_node["meta"],
            "artifact": "03_selection",
            "source": str(ws.shift_cube_path(row["node"]).relative_to(ws.root_dir)),
            "node": row["node"],
            "family": row["family"],
            "feature": row["feature"],
            "params": row["params"],
            "selection": {
                "rank": row["rank"],
                "score": row["score"],
                "method": result["method"]["score"],
            },
        }
        artifact_io.save_selected_node(selected_node, ws.selection_array_path(row), meta)
        copied_npz += 1

        source_xlsx = ws.shift_surface_path(node["id"])
        target_xlsx = ws.selection_surface_path(row)
        try:
            _copy_selected_workbook_sheet(source_xlsx, target_xlsx, int(row["bin"]))
            copied_xlsx += 1
        except PermissionError:
            warnings.append(
                f"selection workbook locked for rank {row['rank']} {row['node']}; "
                "close it in Excel and re-run"
            )
        except FileNotFoundError:
            warnings.append(
                f"selection workbook unavailable for rank {row['rank']} {row['node']}; "
                "run compare to render it"
            )

    ws.write_json(ws.selection_path, result)
    from barrierlab.presentation.selection_plots import write_selected_shift_graphs
    plots = write_selected_shift_graphs(ws, selected)

    report.line(
        f"target ±{abs(delta):.0%} within {horizon}{ws.horizon_unit} · "
        "score = |positive shift − negative shift|"
    )
    report.summary(
        f"wrote {copied_npz} arrays + {copied_xlsx} workbooks + {len(plots)} plots → "
        f"{ws.dir.relative_to(ws.root_dir)}/03_selection"
    )
    report.completed()
    StageReport.warnings(warnings)
