"""
The pipeline's deliverable.

write_barrier_xlsx renders a cube as one tab per condition bin. Each tab is that bin's
whole barrier-by-horizon face, so the workbook holds every value the cube holds -- it is
a faithful view of the measurement, not a summary of it.

Stage 1 workbooks show raw conditional probabilities. Stage 2 shift workbooks show the
same full grid after subtracting the unconditional baseline.
"""

from __future__ import annotations

import re
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from barrierlab.domain.barrier import MIN_BIN_N

_PCT_FMT = '0.0%'
_SHIFT_PCT_FMT = '+0.0%;-0.0%;0.0%'
_PVAL_FMT = '0.0000'

_WHITE, _AMBER, _RED = 'FFFFFF', 'FFD166', 'C00000'
_BLUE = '2A78D6'
_GREEN = '00875A'
_SHIFT_SCALE_LIMIT = 0.30
_FILL_ROW1 = PatternFill('solid', start_color='666666', end_color='666666')
_FILL_ROW2 = PatternFill('solid', start_color='B2B2B2', end_color='B2B2B2')
_FILL_NBAND = PatternFill('solid', start_color='EFEFEF', end_color='EFEFEF')
_FILL_NA = PatternFill('solid', start_color='F7F7F7', end_color='F7F7F7')
_FILL_MID = PatternFill('solid', start_color='DDDDDD', end_color='DDDDDD')
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

_COL_ROW  = 4   # horizon labels
_N_ROW    = 3   # observations behind this bin at each horizon
_DATA_ROW = 5   # first barrier row
_VALID_COL_ROW = 3

_ABBREV = {'rsi', 'ma', 'atr', 'dxy', 'macd', 'bb', 'mvrv', 'wr', 'roc', 'vol'}


def feature_label(node_id: str) -> str:
    words = []
    for p in node_id.split('_'):
        if p.isdigit():
            if words:
                words[-1] += '-' + p
        elif p.lower() in _ABBREV:
            words.append(p.upper())
        else:
            words.append(p.title())
    return ' '.join(words)


def _sheet_names(labels: list[str]) -> list[str]:
    """
    Tab names are the conditions themselves -- 'x < 0.12', '0.12 < x < 0.34' -- so the
    tab strip reads as the ladder of bins with nothing to decode.

    Excel caps a sheet name at 31 chars and forbids : \\ / ? * [ ]. Truncation could in
    principle collide two long labels, so a numeric suffix is appended only when that
    actually happens rather than pre-emptively numbering every tab.
    """
    out: list[str] = []
    for label in labels:
        name = re.sub(r'[:\\/?*\[\]]', '-', label)[:31]
        if name in out:
            base = name[:28]
            k = 2
            while f'{base}~{k}' in out:
                k += 1
            name = f'{base}~{k}'
        out.append(name)
    return out


def _write_headers(ws, row1: str, row2: str, row3: str | None, merge_end: str) -> None:
    rows = [
        (1, row1, Font(bold=True, size=14, color='FFFFFF'), _FILL_ROW1),
        (2, row2, Font(bold=True, size=11, color='000000'), _FILL_ROW2),
    ]
    for r, text, font, fill in rows:
        ws.merge_cells(f'A{r}:{merge_end}{r}')
        c = ws[f'A{r}']
        c.value, c.font, c.fill, c.alignment = text, font, fill, _CENTER
        ws.row_dimensions[r].height = 18


def _condition_title(node_id: str, feature: str, labels: list[str], edges: np.ndarray,
                     bin_index: int, unconditional: bool) -> str:
    display_feature = feature_label(node_id)
    if unconditional:
        condition = 'all'
    elif len(edges) == 0:
        condition = labels[bin_index]
    elif bin_index == 0:
        condition = f'{display_feature} < {float(edges[0]):.3f}'
    elif bin_index == len(edges):
        condition = f'{float(edges[-1]):.3f} < {display_feature}'
    else:
        condition = f'{float(edges[bin_index - 1]):.3f} < {display_feature} < {float(edges[bin_index]):.3f}'
    return f'Condition —— {condition}'


def write_barrier_xlsx(
    cube:    dict,
    path:    Path,
    node_id: str,
    feature: str,
    params:  dict,
    unit:    str = 'd',
) -> None:
    """
    One tab per condition bin; each tab is that bin's full barrier × horizon face.

    Rows run from the highest barrier at the top to the lowest at the bottom, the way a
    price ladder reads: up the sheet is up in price. The zero barrier sits in the middle and is
    shaded, marking the boundary between two different questions -- above it a cell asks
    whether the *high* reached that level, below it whether the *low* did.

    The colour scale is a fixed 0..100% on every tab, so flipping between them shows the
    probability band moving with the condition instead of each tab being rescaled to
    look alike.

    The n band under the header is the observation count behind that bin at each
    horizon; it falls as t grows, because the last t bars have no realized forward
    window.
    """
    conditional_probability = cube['conditional_probability']
    barriers = cube['barriers']
    horizons = cube['horizons']
    bin_observation_counts = cube['bin_observation_counts']
    labels   = cube['meta']['bin_labels']
    bin_edges = cube['bin_edges']
    barrier_count, effective_bin_count, horizon_count = conditional_probability.shape

    order = np.argsort(barriers)[::-1]        # highest barrier on the top row
    end   = get_column_letter(1 + horizon_count)

    wb = Workbook()
    wb.remove(wb.active)

    unconditional = effective_bin_count == 1 and labels[0] == 'all'
    names = _sheet_names(list(labels))

    for b in range(effective_bin_count):
        ws = wb.create_sheet(names[b])
        title = _condition_title(node_id, feature, labels, bin_edges, b, unconditional)
        _write_headers(
            ws,
            title,
            'Conditional barrier-touch probability',
            None,
            merge_end=end,
        )

        c = ws.cell(row=_COL_ROW, column=1, value='barrier')
        c.font, c.alignment = Font(bold=True), _CENTER
        for j, t in enumerate(horizons):
            c = ws.cell(row=_COL_ROW, column=2 + j, value=f'+{int(t)}{unit}')
            c.font, c.alignment = Font(bold=True), _CENTER

        c = ws.cell(row=_N_ROW, column=1, value='n =')
        c.font, c.fill = Font(bold=True, italic=True, size=9), _FILL_NBAND
        for j in range(horizon_count):
            c = ws.cell(row=_N_ROW, column=2 + j, value=int(bin_observation_counts[b, j]))
            c.font, c.fill, c.alignment = Font(italic=True, size=9), _FILL_NBAND, _CENTER

        for r_off, i in enumerate(order):
            r = _DATA_ROW + r_off
            th = float(barriers[i])
            is_zero = abs(th) < 1e-12
            tc = ws.cell(row=r, column=1, value=th)
            tc.number_format = '+0%;-0%;0%'
            tc.font, tc.alignment = Font(bold=True), _CENTER
            if is_zero:
                tc.fill = _FILL_MID
            for j in range(horizon_count):
                v = conditional_probability[i, b, j]
                cell = ws.cell(row=r, column=2 + j,
                               value=None if not np.isfinite(v) else round(float(v), 4))
                cell.number_format = _PCT_FMT
                if not np.isfinite(v):
                    cell.fill = _FILL_NA

        ws.conditional_formatting.add(
            f'B{_DATA_ROW}:{end}{_DATA_ROW + barrier_count - 1}',
            ColorScaleRule(start_type='num', start_value=0,   start_color=_WHITE,
                           mid_type='num',   mid_value=0.5,   mid_color=_AMBER,
                           end_type='num',   end_value=1.0,   end_color=_RED))

        ws.column_dimensions['A'].width = 8
        for j in range(horizon_count):
            ws.column_dimensions[get_column_letter(2 + j)].width = 6.5
        ws.freeze_panes = f'B{_DATA_ROW}'

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_shift_xlsx(
    cube:    dict,
    path:    Path,
    node_id: str,
    feature: str,
    params:  dict,
    unit:    str = 'd',
) -> None:
    """
    One tab per condition bin; each tab is the full baseline-subtracted surface.

    Values are probability differences from the baseline, displayed as percentages. The color scale is
    centered at zero: blue means the barrier is touched less often than unconditional,
    red means more often.
    """
    values = cube['probability_shift']
    barriers = cube['barriers']
    horizons = cube['horizons']
    bin_observation_counts = cube['bin_observation_counts']
    labels   = cube['meta']['bin_labels']
    bin_edges = cube['bin_edges']
    barrier_count, effective_bin_count, horizon_count = values.shape

    order = np.argsort(barriers)[::-1]
    end = get_column_letter(1 + horizon_count)
    lim = _SHIFT_SCALE_LIMIT

    wb = Workbook()
    wb.remove(wb.active)

    unconditional = effective_bin_count == 1 and labels[0] == 'all'
    names = _sheet_names(list(labels))

    for b in range(effective_bin_count):
        ws = wb.create_sheet(names[b])
        title = _condition_title(node_id, feature, labels, bin_edges, b, unconditional)
        _write_headers(
            ws,
            title,
            'Conditional probability minus baseline probability',
            None,
            merge_end=end,
        )

        c = ws.cell(row=_COL_ROW, column=1, value='barrier')
        c.font, c.alignment = Font(bold=True), _CENTER
        for j, t in enumerate(horizons):
            c = ws.cell(row=_COL_ROW, column=2 + j, value=f'+{int(t)}{unit}')
            c.font, c.alignment = Font(bold=True), _CENTER

        c = ws.cell(row=_N_ROW, column=1, value='n =')
        c.font, c.fill = Font(bold=True, italic=True, size=9), _FILL_NBAND
        for j in range(horizon_count):
            c = ws.cell(row=_N_ROW, column=2 + j, value=int(bin_observation_counts[b, j]))
            c.font, c.fill, c.alignment = Font(italic=True, size=9), _FILL_NBAND, _CENTER

        for r_off, i in enumerate(order):
            r = _DATA_ROW + r_off
            th = float(barriers[i])
            is_zero = abs(th) < 1e-12
            tc = ws.cell(row=r, column=1, value=th)
            tc.number_format = '+0%;-0%;0%'
            tc.font, tc.alignment = Font(bold=True), _CENTER
            if is_zero:
                tc.fill = _FILL_MID
            for j in range(horizon_count):
                v = values[i, b, j]
                cell = ws.cell(row=r, column=2 + j,
                               value=None if not np.isfinite(v) else float(v))
                cell.number_format = _SHIFT_PCT_FMT
                if not np.isfinite(v):
                    cell.fill = _FILL_NA

        rng = f'B{_DATA_ROW}:{end}{_DATA_ROW + barrier_count - 1}'
        ws.conditional_formatting.add(
            rng,
            ColorScaleRule(start_type='num', start_value=-lim, start_color=_BLUE,
                           mid_type='num', mid_value=0, mid_color=_WHITE,
                           end_type='num', end_value=lim, end_color=_RED))

        ws.column_dimensions['A'].width = 8
        for j in range(horizon_count):
            ws.column_dimensions[get_column_letter(2 + j)].width = 8.5
        ws.freeze_panes = f'B{_DATA_ROW}'

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
