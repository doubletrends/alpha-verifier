"""Stage 3 command: select the strongest whole-node Bayes predictors."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

from domain import bayes, selection, shift
from infrastructure.workspaces.workspace import BASELINE_NODE, Workspace


def _copy_selected_workbook_sheet(source: Path, target: Path, bin_index: int) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    wb = load_workbook(target)
    keep = wb.sheetnames[int(bin_index)]
    for sheet_name in list(wb.sheetnames):
        if sheet_name != keep:
            del wb[sheet_name]
    wb.save(target)
    wb.close()


def cmd_selection(ws: Workspace) -> None:
    """Stage 3: score every node's full bin table and retain top distinct nodes."""
    nodes = [
        n for n in ws.catalog.all_nodes()
        if n["id"] != BASELINE_NODE and ws.has_shift_cube(n["family"], n["id"])
    ]
    if not nodes:
        print("No shift arrays to select from - run shift first.")
        return

    def load_cube(node: dict) -> dict:
        return shift.load(ws.shift_cube_path(node["family"], node["id"]))

    baseline_path = ws.shift_cube_path("_base", BASELINE_NODE)
    if not baseline_path.exists():
        print("No baseline shift array - run shift first.")
        return
    delta, horizon = ws.bayes_target(shift.load(baseline_path)["base"])
    result = selection.rank_nodes(
        nodes,
        load_cube,
        ws.selection_top_k,
        delta,
        horizon,
        bayes.SHRINK_K,
    )
    result["workspace"] = ws.dir.name
    result["families"] = sorted({n["family"] for n in nodes})

    selected = result["selected"]
    copied_npz = 0
    copied_xlsx = 0
    by_id = {n["id"]: n for n in nodes}
    for row in selected:
        node = by_id[row["node"]]
        source = load_cube(node)
        selected_node = selection.selected_node_from_shift_cube(source, row)
        meta = {
            **selected_node["meta"],
            "artifact": "03_selection",
            "source": str(ws.shift_cube_path(row["family"], row["node"]).relative_to(ws.root_dir)),
            "node": row["node"],
            "family": row["family"],
            "feature": row["feature"],
            "params": row["params"],
            "selection": {
                "rank": row["rank"],
                "score": row["score"],
                "method": result["method"]["score"],
            },
        }
        selection.save_selected_node(selected_node, ws.selection_array_path(row), meta)
        copied_npz += 1

        source_xlsx = ws.shift_surface_path(node["family"], node["id"])
        target_xlsx = ws.selection_surface_path(row)
        try:
            _copy_selected_workbook_sheet(source_xlsx, target_xlsx, int(row["bin"]))
            copied_xlsx += 1
        except PermissionError:
            print(f"  rank {row['rank']:>3} {row['node']:<26} [locked] close it in Excel and re-run")
        except FileNotFoundError:
            print(f"  rank {row['rank']:>3} {row['node']:<26} [no shift workbook] run shift to render it")

    ws.write_json(ws.selection_path, result)

    print(f"\n=== 3. Selection [{ws.dir.name}] ===")
    print(
        f"  scored {len(result['candidates'])} whole nodes from 02_shift; "
        f"selected top {len(selected)} distinct nodes by conditional information"
    )
    print(
        f"  target = P(touch {delta:+.0%} within {horizon}{ws.horizon_unit})  |  "
        "score = weighted information across all bins\n"
    )
    if selected:
        print(f"  {'rank':>4}  {'node':<26}{'family':<13}{'bin':>5}  {'bits':>7}  {'cover':>5}  representative bin")
        print(f"  {'-'*4}  {'-'*26}{'-'*13}{'-'*5}  {'-'*7}  {'-'*5}  {'-'*37}")
        for row in selected:
            c = row["best_cell"]
            print(
                f"  {row['rank']:>4}  {row['node']:<26}{row['family']:<13}"
                f"{row['bin_number']:>5}  {row['score_bits']:>7.4f}  {row['effective_bins']:>5.1f}  "
                f"dev={c['dev']:+.1f}pp n={c['bin_n']}"
            )

    print(f"\n  wrote {copied_npz} .npz artifacts under {ws.dir.relative_to(ws.root_dir)}/03_selection/")
    print(f"  wrote {copied_xlsx} workbooks under {ws.dir.relative_to(ws.root_dir)}/03_selection/")
