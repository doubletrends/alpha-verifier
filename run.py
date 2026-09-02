"""Agent run entry point.

Every command runs against one barrier — the outcome declared in the workspace's
universe.json, or overridden with --outcome / --threshold. All output is scoped by
that barrier's event label (dd10, run10, ...), so two barriers never collide.

Usage:
    python run.py --workspace btc_daily_14days --family rsi
    python run.py --workspace btc_daily_14days --node rsi_14
    python run.py --workspace btc_daily_14days --next
    python run.py --workspace btc_daily_14days --status
    python run.py --workspace btc_daily_14days --read rsi_14
    python run.py --workspace btc_daily_14days --probe
    python run.py --workspace btc_daily_14days --skew
    python run.py --workspace btc_daily_14days --outcome runup --family rsi
    python run.py --workspace btc_daily_14days --validate
"""

import argparse
import json
import sys
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from data import features, fetcher
from engine import matrix as engine
from engine import writer
from engine import combiner as cmb
from engine import outcomes
from engine import validate as val
from engine import charts
from tree.tree import load_tree, all_nodes, find_node, all_in_family
from workspace import Workspace


def _h_num(label: str) -> int:
    """Extract the numeric part from a horizon label like '+7d' or '+12h'."""
    return int(''.join(c for c in label if c.isdigit()))


def _tested_by_family(
    ws: Workspace,
    families_filter: list[str] | None = None,
) -> dict[str, list[dict]]:
    """
    Tested nodes grouped by family, for the workspace's current event.

    'Tested' means the node's xlsx exists under this event's directory — status is
    read from disk, not from universe.json, so each barrier tracks its own progress.
    """
    tree     = load_tree(ws.tree_path)
    by_family: dict[str, list[dict]] = defaultdict(list)
    for node in all_nodes(tree):
        if ws.is_tested(node['family'], node['id']):
            by_family[node['family']].append(node)
    if families_filter:
        by_family = {f: ns for f, ns in by_family.items() if f in families_filter}
    return by_family


def _nodes_todo(ws: Workspace, family: str | None = None) -> list[dict]:
    """Nodes not yet tested and not previously skipped, for the current event."""
    tree  = load_tree(ws.tree_path)
    nodes = all_in_family(tree, family) if family else all_nodes(tree)
    return [n for n in nodes if ws.node_status(n['family'], n['id']) == 'pending']


def _select_best_nodes(
    ws:        Workspace,
    by_family: dict[str, list[dict]],
) -> dict[str, tuple[dict, float]]:
    """
    For each family, pick the node with the highest peak absolute deviation
    at the pivot horizon (n ≥ 30 slices required).

    Returns {family: (node, peak_signal)}.
    """
    best: dict[str, tuple[dict, float]] = {}
    for fam, nodes in by_family.items():
        for node in nodes:
            path = ws.output_path(fam, node['id'])
            if not path.exists():
                continue
            try:
                fn_tbl = cmb.load_fn_table(path)
                pk     = cmb.peak_signal(fn_tbl, ws.pivot_horizon, min_n=30)
                if fam not in best or pk > best[fam][1]:
                    best[fam] = (node, pk)
            except Exception:
                pass
    return best


def _make_loader(ws: Workspace):
    """Return a cached fetch function: load(sources) → DataFrame, fetched once per unique source set."""
    cache: dict[tuple, pd.DataFrame] = {}
    def load(sources: list) -> pd.DataFrame:
        key = tuple(sorted(sources))
        if key not in cache:
            cache[key] = fetcher.fetch(list(sources), start=ws.start_date, asset=ws.asset).dropna(subset=['close'])
        return cache[key]
    return load


# ── core run ──────────────────────────────────────────────────────────────────

def run_node(ws: Workspace, node_id: str, regen: bool = False) -> None:
    tree = load_tree(ws.tree_path)
    node = find_node(tree, node_id)

    family  = node['family']
    out_path = ws.output_path(family, node_id)

    print(f"\n{'=' * 60}")
    print(f"Node     : {node['id']}")
    print(f"Family   : {family}  |  Category: {node['category']}")
    print(f"Feature  : {node['feature']}  params={node['params']}")
    print(f"Outcome  : {ws.outcome}  ->  P({ws.outcome_expr})")
    print(f"Sources  : {node['data']}")
    if node.get('derived_from'):
        print(f"Derived  : from {node['derived_from']}")
    print(f"{'=' * 60}\n")

    print('Fetching data...')
    data = fetcher.fetch(node['data'], start=ws.start_date, asset=ws.asset).dropna(subset=['close'])
    if data.empty:
        print('Fetch returned empty DataFrame — skipping.')
        return
    print(f'Rows: {len(data)}  ({data.index[0].date()} → {data.index[-1].date()})')

    print(f"Computing feature '{node['feature']}'...")
    feat = features.compute(data, node['feature'], node['params']).reindex(data.index)

    n_valid = int(feat.notna().sum())
    if n_valid < ws.min_obs:
        print(f'Only {n_valid} valid observations — skipping.')
        if not regen:
            ws.record_skip(family, node_id, f'only {n_valid} valid observations (min {ws.min_obs})')
        return

    lo         = float(np.nanpercentile(feat.dropna(), 2))
    hi         = float(np.nanpercentile(feat.dropna(), 98))
    thresholds = np.linspace(lo, hi, ws.n_thresholds)
    print(f'Feature p2→p98: [{lo:.4g}, {hi:.4g}]')

    base_rate = engine.compute_base_rate(
        data, ws.horizons, horizon_unit=ws.horizon_unit,
        outcome=ws.outcome, outcome_params=ws.outcome_params,
    )

    # A barrier that almost never gets touched cannot be measured by conditional
    # counting: every threshold slice holds a handful of events and each deviation
    # is noise. θ has to be scaled to the asset and horizon, so fail loudly rather
    # than emit a surface that looks real. (10% over 14 BTC days is a 21.7% event;
    # the same 10% over 24 NASDAQ hours is a 0.4% one.)
    p0 = float(base_rate['base_rate'].dropna().max() or 0.0) / 100.0
    if p0 < outcomes.MIN_BASE_RATE:
        msg = (f'base rate {p0:.2%} < {outcomes.MIN_BASE_RATE:.0%} — barrier too rare '
               f'to measure at this threshold')
        print(f'{msg}. Lower --threshold for this workspace.')
        if not regen:
            ws.record_skip(family, node_id, msg)
        return
    p_below, p_above = engine.compute_matrix(
        feat, thresholds, ws.horizons, data, horizon_unit=ws.horizon_unit,
        outcome=ws.outcome, outcome_params=ws.outcome_params,
    )

    writer.write_xlsx(
        p_below, p_above, base_rate, out_path, node_id,
        event=ws.event, title=ws.outcome_spec['title'], expr=ws.outcome_expr,
        sign=ws.outcome_sign,
    )

    br   = {row: round(float(base_rate.loc[row, 'base_rate']), 2) for row in base_rate.index}
    spot = '  '.join(f'{h}={br[h]:.1f}%' for h in ws.display_horizons if h in br)
    print(f'Base rate P({ws.outcome_expr}): {spot}')

    if not regen:
        # The node ran, so clear any stale skip recorded for this event.
        ws.clear_skip(family, node_id)
        log_entry = {
            'node_id':      node_id,
            'family':       family,
            'category':     node['category'],
            'feature':      node['feature'],
            'params':       node['params'],
            'derived_from': node.get('derived_from'),
            'timestamp':    datetime.now(timezone.utc).isoformat(),
            'matrix_path':  str(out_path.relative_to(ROOT)),
            'outcome':      ws.outcome,
            'event':        ws.event,
            'n_obs':        n_valid,
            'horizons':     ws.horizons,
            'n_thresholds': ws.n_thresholds,
            'base_rate':    br,
        }
        log_path = ws.family_log(family)
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(log_path, 'a', encoding='utf-8') as f:
            f.write(json.dumps(log_entry) + '\n')

    print(f"\nDone → {out_path.relative_to(ROOT)}")


# ── CLI commands ──────────────────────────────────────────────────────────────

def cmd_family(ws: Workspace, family_name: str, rerun: bool = False) -> None:
    tree = load_tree(ws.tree_path)

    if rerun:
        todo, label = all_in_family(tree, family_name), 'node'
    else:
        todo, label = _nodes_todo(ws, family_name), 'pending node'

    if not todo:
        families = sorted({n['family'] for n in all_nodes(tree)})
        if all_in_family(tree, family_name):
            print(f"No {label}s in family '{family_name}' for event '{ws.event}' "
                  f"(all tested or skipped). Use --rerun to run them anyway.")
        else:
            print(f"No family '{family_name}'. Available: {', '.join(families)}")
        return

    print(f"Running {len(todo)} {label}(s) in family '{family_name}' for event '{ws.event}'...")
    for node in todo:
        run_node(ws, node['id'])


def cmd_next(ws: Workspace) -> None:
    todo = _nodes_todo(ws)
    if not todo:
        print(f"All nodes tested or skipped for event '{ws.event}'.")
        return
    run_node(ws, todo[0]['id'])


def cmd_list(ws: Workspace) -> None:
    by_fam = defaultdict(list)
    for n in _nodes_todo(ws):
        by_fam[n['family']].append(n['id'])
    if not by_fam:
        print(f"No pending nodes for event '{ws.event}'.")
        return
    print(f"\nPending nodes by family  [event: {ws.event}]:")
    for fam in sorted(by_fam):
        print(f"  {fam:<20s}  {', '.join(by_fam[fam])}")


def cmd_status(ws: Workspace) -> None:
    """Progress for the current event, read off the filesystem."""
    tree   = load_tree(ws.tree_path)
    by_fam = defaultdict(lambda: {'pending': 0, 'tested': 0, 'skipped': 0})
    for n in all_nodes(tree):
        by_fam[n['family']][ws.node_status(n['family'], n['id'])] += 1

    print(f"\n=== Status [{ws.dir.name}] — event '{ws.event}': P({ws.outcome_expr}) ===")
    print(f"\n  {'family':<18}  {'pending':>7}  {'tested':>6}  {'skipped':>7}")
    print('  ' + '-' * 44)
    for fam in sorted(by_fam):
        c = by_fam[fam]
        print(f"  {fam:<18s}  {c['pending']:>7}  {c['tested']:>6}  {c['skipped']:>7}")
    totals = {k: sum(by_fam[f][k] for f in by_fam) for k in ('pending', 'tested', 'skipped')}
    print('  ' + '-' * 44)
    print(f"  {'TOTAL':<18s}  {totals['pending']:>7}  {totals['tested']:>6}  {totals['skipped']:>7}")

    skips = [(fam, nid, r) for fam in sorted(by_fam)
             for nid, r in ws.load_skips(fam).items()]
    if skips:
        print()
        for fam, nid, reason in skips:
            print(f"  skipped: {nid:<24} ({fam}) — {reason}")


def _significant_rows(wb, read_h: list[str], min_n: int, min_dev: float):
    """Yield (kind, cond, n_raw, devs) for every significant row across ABOVE and BELOW sheets."""
    for sheet_name in ('above', 'below'):
        ws_sheet = writer.find_sheet(wb, sheet_name)
        rows = list(ws_sheet.iter_rows(values_only=True))
        hdrs = rows[4]
        ci   = {h: i for i, h in enumerate(hdrs) if h is not None}
        for row in rows[5:]:
            cond  = row[ci['condition']]
            n_raw = row[ci['n']]
            if not n_raw:
                continue
            if int(n_raw) < min_n:
                continue
            devs = {h: row[ci[h]] for h in read_h if h in ci and row[ci[h]] is not None}
            if any(abs(v) >= min_dev for v in devs.values()):
                yield sheet_name, cond, n_raw, devs


def cmd_read(ws: Workspace, node_id: str) -> None:
    tree = load_tree(ws.tree_path)
    node = find_node(tree, node_id)
    path = ws.output_path(node['family'], node_id)

    if not path.exists():
        print(f"No output file: {path}")
        return

    wb = openpyxl.load_workbook(path, data_only=True)
    read_h = ws.display_horizons

    print(f"\n{node_id}  [{node['family']}]  {node['feature']}  params={node['params']}")

    by_sheet = {'above': [], 'below': []}
    for kind, cond, n_raw, devs in _significant_rows(wb, read_h, ws.read_min_n, ws.read_min_dev):
        by_sheet[kind].append((cond, n_raw, devs))

    h_hdr = ''.join(f'{h:>8}' for h in read_h)
    h_sep = ''.join('--------' for _ in read_h)
    for direction in ('above', 'below'):
        significant = by_sheet[direction]
        if not significant:
            print(f"\n  [{direction}]  no rows ≥ {ws.read_min_dev}pp with n ≥ {ws.read_min_n}")
            continue
        print(f"\n  [{direction}]")
        print(f"  {'condition':<26}  {'n':>12}  {h_hdr}")
        print(f"  {'-'*26}  {'-'*12}  {h_sep}")
        for cond, n_raw, devs in significant:
            def fmt(h):
                v = devs.get(h)
                return f'{v:+.1f}pp' if v is not None else '   n/a'
            h_vals = ''.join(f'{fmt(h):>8}' for h in read_h)
            print(f"  {str(cond):<26}  {str(n_raw):>12}  {h_vals}")


def cmd_probe(ws: Workspace, families_filter: list[str] | None = None) -> None:
    by_family = _tested_by_family(ws, families_filter)
    best      = _select_best_nodes(ws, by_family)
    get_data  = _make_loader(ws)

    print('Fetching latest data...')
    ohlcv_data = get_data(['ohlcv'])
    base_rate  = engine.compute_base_rate(
        ohlcv_data, ws.horizons, horizon_unit=ws.horizon_unit,
        outcome=ws.outcome, outcome_params=ws.outcome_params,
    )
    br_series  = base_rate['base_rate']

    key_h  = ws.key_horizons

    print(f"\n=== Signal Probe [{ws.dir.name}] — {date.today()} ===")
    print(f"\n  One node per family, auto-selected by peak |{ws.pivot_horizon}| edge (n≥30 slices):\n")
    _h_hdrs = ''.join(f'{h:>8}' for h in key_h)
    _h_seps = ''.join('--------' for _ in key_h)
    print(f"  {'family':<18}  {'node':<26}  {'current x':>11}  {'n':>6}  {'weight':>6}  {_h_hdrs}")
    print(f"  {'-'*18}  {'-'*26}  {'-'*11}  {'------':>6}  {'------':>6}  {_h_seps}")

    contributions = []
    weights       = {}

    for fam in sorted(best):
        node, _ = best[fam]
        path    = ws.output_path(fam, node['id'])
        try:
            data = get_data(node['data'])
            if data.empty:
                raise ValueError('empty data')
            feat      = features.compute(data, node['feature'], node['params']).reindex(data.index)
            valid     = feat.dropna()
            if valid.empty:
                raise ValueError('no valid feature values')
            current_x = float(valid.iloc[-1])

            fn_tbl = cmb.load_fn_table(path)
            devs   = cmb.eval_at(fn_tbl, current_x)
            n_val  = int(devs['n']) if pd.notna(devs.get('n', np.nan)) else 0
            w      = cmb.shrink_weight(n_val)
            devs_h = devs[[c for c in devs.index if c.startswith('+')]]
            contributions.append((node['id'], devs_h))
            weights[node['id']] = w

            def fmtd(h: str) -> str:
                v = devs_h[h] if h in devs_h.index else np.nan
                return f'{float(v):+.1f}pp' if pd.notna(v) else '   n/a'

            _h_vals = ''.join(f'{fmtd(h):>8}' for h in key_h)
            print(f"  {fam:<18}  {node['id']:<26}  {current_x:>11.4g}  {n_val:>6}  {w:>6.2f}  {_h_vals}")

        except Exception as e:
            print(f"  {fam:<18}  {node['id']:<26}  {'[error]':>11}  — {e}")

    if not contributions:
        print('\nNo contributions computed.')
        return

    result = cmb.combine(contributions, br_series, horizons=key_h, weights=weights)

    n_sig = len(contributions)
    print(f"\n  Naive Bayes combination [{n_sig} signals, one per family]:\n")
    print(f"  {'horizon':<8}  {'base rate':>10}  {'combined':>10}  {'edge':>10}")
    print(f"  {'-'*8}  {'-'*10}  {'-'*10}  {'-'*10}")
    for h in key_h:
        br   = float(br_series.loc[h]) if h in br_series.index else np.nan
        comb = float(result.loc[h, 'combined']) if h in result.index else np.nan
        edge = float(result.loc[h, 'edge'])     if h in result.index else np.nan
        print(f"  {h:<8}  {br:>9.1f}%  {comb:>9.1f}%  {edge:>+10.1f}pp")

    print()
    print('  * Assumes feature independence. Correlated signals may inflate the estimate.')
    print()




def cmd_validate(
    ws:              Workspace,
    families_filter: list[str] | None = None,
    n_shifts:        int = val.DEFAULT_SHIFTS,
) -> None:
    """
    Test every tested node against its circular-shift null and write validation.json.

    A node's peak deviation is a max over ~30 bins, which is biased upward even under
    pure noise; this reports how far the real peak sits inside the null that the same
    node would have produced by chance. Bonferroni correction is applied across every
    test in the sweep, so a 'structure' verdict already accounts for the search size.
    """
    by_family = _tested_by_family(ws, families_filter)
    get_data  = _make_loader(ws)
    nodes     = [n for fam in sorted(by_family) for n in by_family[fam]]

    if not nodes:
        print('No tested nodes to validate.')
        return

    read_h  = ws.display_horizons
    n_tests = len(nodes) * len(read_h)

    print()
    print(f"=== Validation [{ws.dir.name}] - outcome '{ws.outcome}' ===")
    print(f"  event   : P({ws.outcome_expr})")
    print(f"  sweep   : {len(nodes)} nodes x {len(read_h)} horizons = {n_tests} tests, "
          f"{n_shifts} circular shifts each")
    print(f"  alpha   : 0.05 / {n_tests} = {0.05 / n_tests:.2e} (Bonferroni)")
    print()

    h_hdr = ''.join(f'{h:>28}' for h in read_h)
    print(f"  {'node':<26}  {'family':<12}{h_hdr}")
    print(f"  {'-' * 26}  {'-' * 12}{''.join('-' * 28 for _ in read_h)}")

    results = {}
    for node in nodes:
        fam = node['family']
        try:
            data  = get_data(node['data'])
            feat  = features.compute(data, node['feature'], node['params']).reindex(data.index)
            valid = feat.dropna()
            if valid.empty:
                raise ValueError('no valid feature values')
            lo, hi     = np.nanpercentile(valid, 2), np.nanpercentile(valid, 98)
            thresholds = np.linspace(float(lo), float(hi), ws.n_thresholds)
        except Exception as e:
            print(f"  {node['id']:<26}  {fam:<12}  [skip] {e}")
            continue

        cells, per_h = [], {}
        for h in read_h:
            ev  = outcomes.compute(data, ws.outcome, _h_num(h), ws.outcome_params)
            res = val.null_test(feat, ev, thresholds, n_shifts=n_shifts)
            res['verdict'] = val.verdict(res['p_value'], n_tests=n_tests,
                                         p_floor=res.get('p_floor', 0.0))
            per_h[h] = res
            if pd.isna(res['real']):
                cells.append(f"{'insufficient':>28}")
            else:
                mark = {'structure': ' **', 'nominal': ' * ', 'underpowered': ' ? ',
                        'noise': '   ', 'insufficient': '   '}[res['verdict']]
                cells.append(f"{res['real']:>9.1f} vs{res['null_p95']:>6.1f}  p={res['p_value']:<5.3f}{mark}")

        results[node['id']] = {'family': fam, 'horizons': per_h}
        print(f"  {node['id']:<26}  {fam:<12}{''.join(cells)}")

    print()
    print('  columns: real peak  vs  null p95   p-value')
    print('  ** clears Bonferroni   * nominal only   ? at the resolution floor, needs more shifts')
    print()

    # Merge into any prior sweep for this outcome, so validating one family at a
    # time never discards the rest. Bonferroni is then applied over every test in
    # the merged file — narrowing --families must not quietly shrink the correction
    # for a search that was already run wide.
    path = ws.validation_path

    merged = {}
    if path.exists():
        try:
            with open(path, encoding='utf-8') as f:
                prior = json.load(f)
            if prior.get('outcome', {}).get('name') == ws.outcome:
                merged = prior.get('nodes', {})
        except (json.JSONDecodeError, OSError):
            pass
    merged.update(results)

    n_merged = sum(len(n['horizons']) for n in merged.values())
    for node_res in merged.values():
        for res in node_res['horizons'].values():
            res['verdict'] = val.verdict(res['p_value'], n_tests=n_merged,
                                         p_floor=res.get('p_floor', 0.0))

    results   = merged
    n_tests   = n_merged
    flat      = [r for n in results.values() for r in n['horizons'].values()]
    n_struct  = sum(1 for r in flat if r['verdict'] == 'structure')
    n_under   = sum(1 for r in flat if r['verdict'] == 'underpowered')
    n_nominal = sum(1 for r in flat if r['verdict'] == 'nominal')
    expected  = 0.05 * len(flat)
    print(f"  (Bonferroni over {n_tests} tests in {path.name}, including earlier runs)")
    print(f"  clears Bonferroni : {n_struct:>4} / {len(flat)}")
    print(f"  at floor (needs +): {n_under:>4} / {len(flat)}")
    print(f"  nominal (p<0.05)  : {n_nominal:>4} / {len(flat)}   (expected by chance ~{expected:.1f})")
    if n_under:
        need = val.shifts_needed(n_tests=n_tests)
        print(f"  -> re-run with --shifts {need} to resolve the {n_under} sitting at the floor.")
    if n_struct == 0 and n_under == 0 and n_nominal <= expected:
        print('  -> no more hits than chance alone would produce; treat as no structure.')

    out = {
        'workspace':  ws.dir.name,
        'outcome':    {'name': ws.outcome, 'params': ws.outcome_params, 'expr': ws.outcome_expr},
        'generated':  datetime.now(timezone.utc).isoformat(),
        # A merged file can hold records from runs of differing depth, so report the
        # range actually present rather than just this run's setting.
        'n_shifts':   {'min': min(r.get('n_shifts', n_shifts) for r in flat),
                       'max': max(r.get('n_shifts', n_shifts) for r in flat),
                       'last_run': n_shifts},
        'n_tests':    n_tests,
        'alpha_bonf': 0.05 / n_tests,
        'summary':    {'structure': n_struct, 'underpowered': n_under,
                       'nominal': n_nominal, 'total': len(flat)},
        'nodes':      results,
    }
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(out, f, indent=2)
    print()
    print(f"Wrote {path.relative_to(ROOT)}")


def cmd_skew(ws: Workspace, families_filter: list[str] | None = None) -> None:
    """
    Measure both barriers on identical bars and report the asymmetry.

    A condition that raises P(touch -theta) might be picking out genuine downside
    risk, or might simply be picking out volatility -- in which case it raises
    P(touch +theta) by just as much and carries no directional information at all.
    Counting one barrier alone cannot tell those apart. This runs the same feature,
    the same bars and the same thresholds against both and prints the difference.

    The baseline is itself asymmetric -- BTC touches +10% more often than -10%
    simply because it drifts up -- so the skew reported is a deviation-vs-deviation
    comparison, each barrier already net of its own base rate.
    """
    mirror = ws.mirror_outcome
    if mirror is None:
        print(f"Outcome '{ws.outcome}' has no mirror barrier; --skew needs a pair.")
        return

    tree     = load_tree(ws.tree_path)
    get_data = _make_loader(ws)
    nodes    = [n for n in all_nodes(tree)
                if not families_filter or n['family'] in families_filter]
    if not nodes:
        print('No nodes selected.')
        return

    read_h = ws.display_horizons
    down, up_ = (ws.outcome, mirror) if ws.outcome_sign < 0 else (mirror, ws.outcome)
    d_params = dict(ws.outcome_params)
    u_params = dict(ws.outcome_params)

    print()
    print(f"=== Barrier skew [{ws.dir.name}] ===")
    print(f"  down : P({outcomes.describe(down, d_params)})")
    print(f"  up   : P({outcomes.describe(up_, u_params)})")
    print(f"  skew = (down deviation) - (up deviation), both in pp from their own base rate.")
    print(f"  positive skew = the condition tilts downside beyond what it does to upside.")
    print()

    hdr = ''.join(f'{h:>26}' for h in read_h)
    print(f"  {'node':<26}  {'family':<12}{hdr}")
    print(f"  {'-'*26}  {'-'*12}{''.join('-'*26 for _ in read_h)}")

    rows = []
    for node in nodes:
        fam = node['family']
        try:
            data = get_data(node['data'])
            feat = features.compute(data, node['feature'], node['params']).reindex(data.index)
            valid = feat.dropna()
            if valid.empty:
                raise ValueError('no valid feature values')
            lo, hi = np.nanpercentile(valid, 2), np.nanpercentile(valid, 98)
            thresholds = np.linspace(float(lo), float(hi), ws.n_thresholds)
        except Exception as e:
            print(f"  {node['id']:<26}  {fam:<12}  [skip] {e}")
            continue

        cells, per_h = [], {}
        for h in read_h:
            h_n = _h_num(h)
            ev_d = outcomes.compute(data, down, h_n, d_params)
            ev_u = outcomes.compute(data, up_,  h_n, u_params)
            # Peak deviation for each barrier, over the same bins on the same bars.
            pk_d, x_d = _peak_dev_at(feat, ev_d, thresholds)
            pk_u, _   = _peak_dev_at(feat, ev_u, thresholds, at_bin=x_d)
            if pk_d is None:
                cells.append(f"{'--':>26}")
                continue
            skew = pk_d - pk_u
            per_h[h] = {'down_dev': round(pk_d, 2), 'up_dev': round(pk_u, 2),
                        'skew': round(skew, 2)}
            cells.append(f"{pk_d:>+7.1f}/{pk_u:>+6.1f} -> {skew:>+6.1f}")
        if per_h:
            rows.append((node['id'], fam, per_h))
        print(f"  {node['id']:<26}  {fam:<12}{''.join(cells)}")

    print()
    print('  columns: down dev / up dev -> skew, all in pp, at the bin where |down dev| peaks')
    print()

    pivot = read_h[len(read_h) // 2]
    ranked = sorted((r for r in rows if pivot in r[2]),
                    key=lambda r: -abs(r[2][pivot]['skew']))[:8]
    if ranked:
        print(f"  Strongest asymmetry at {pivot}:")
        for nid, fam, per_h in ranked:
            rec = per_h[pivot]
            kind = ('downside-specific' if rec['skew'] > 5 else
                    'upside-specific'   if rec['skew'] < -5 else
                    'symmetric (volatility, not direction)')
            print(f"    {nid:<26} {fam:<12} skew {rec['skew']:>+6.1f}pp   {kind}")
        print()


def _peak_dev_at(feature, outcome, thresholds, at_bin=None):
    """
    Peak |deviation from base rate| across threshold bins, in pp.

    Returns (deviation, bin_index). When at_bin is given, the deviation is read at
    that bin instead of at its own peak -- which is what makes the two barriers
    comparable: both are evaluated on the *same* slice of feature space.
    """
    aligned = pd.concat([feature.rename('x'), outcome.rename('y')], axis=1).dropna()
    if len(aligned) < 100:
        return None, None
    x = aligned['x'].to_numpy(dtype=float)
    y = aligned['y'].to_numpy(dtype=float)
    idx    = np.searchsorted(thresholds, x)
    n_bins = len(thresholds) + 1
    counts = np.bincount(idx, minlength=n_bins)
    sums   = np.bincount(idx, weights=y, minlength=n_bins)
    ok     = counts >= 30
    if not ok.any():
        return None, None
    rates = np.full(n_bins, np.nan)
    rates[ok] = sums[ok] / counts[ok]
    devs = (rates - y.mean()) * 100.0
    if at_bin is not None:
        return (float(devs[at_bin]) if not np.isnan(devs[at_bin]) else 0.0), at_bin
    b = int(np.nanargmax(np.abs(devs)))
    return float(devs[b]), b


def cmd_backtest(ws: Workspace, families_filter: list[str] | None = None) -> None:
    by_family = _tested_by_family(ws, families_filter)
    best      = _select_best_nodes(ws, by_family)
    get_data  = _make_loader(ws)

    print('Fetching data and computing feature series...')
    ohlcv_data = get_data(['ohlcv'])
    base_rate  = engine.compute_base_rate(
        ohlcv_data, ws.horizons, horizon_unit=ws.horizon_unit,
        outcome=ws.outcome, outcome_params=ws.outcome_params,
    )
    br_series  = base_rate['base_rate']

    selected: list[dict]               = []
    feat_map:  dict[str, pd.Series]    = {}
    fn_map:    dict[str, pd.DataFrame] = {}

    for fam in sorted(best):
        node, _ = best[fam]
        path    = ws.output_path(fam, node['id'])
        try:
            data = get_data(node['data'])
            if data.empty:
                continue
            feat = features.compute(data, node['feature'], node['params']).reindex(data.index)
            fn_map[node['id']]   = cmb.load_fn_table(path)
            feat_map[node['id']] = feat
            selected.append(node)
            print(f"  loaded  {node['id']}  ({fam})")
        except Exception as e:
            print(f"  [skip]  {node['id']}: {e}")

    if not selected:
        print('No nodes loaded.')
        return

    close  = ohlcv_data['close']
    # Realizations come from the configured barrier. This used to be hard-coded to
    # close[t+h] > close[t], which silently scored every non-directional outcome
    # against the wrong event.
    events = {h: outcomes.compute(ohlcv_data, ws.outcome, h, ws.outcome_params)
              for h in ws.horizons}
    min_date = close.index[250]
    max_date = close.index[-(ws.horizons[-1] + 1)]
    sample_dates = []
    for dt in pd.date_range(start=min_date, end=max_date, freq=ws.sample_freq):
        pos = close.index.searchsorted(dt)
        if pos < len(close.index):
            sample_dates.append(close.index[pos])

    print(f'\nRunning combiner across {len(sample_dates)} sample dates...')

    key_h   = ws.key_horizons
    u       = ws.horizon_unit
    results = []

    for dt in sample_dates:
        pos = close.index.get_loc(dt)

        contribs = []
        for node in selected:
            fs = feat_map[node['id']]
            if dt not in fs.index or pd.isna(fs.loc[dt]):
                continue
            devs   = cmb.eval_at(fn_map[node['id']], float(fs.loc[dt]))
            devs_h = devs[[c for c in devs.index if c.startswith('+')]]
            contribs.append((node['id'], devs_h))

        if not contribs:
            continue

        result = cmb.combine(contribs, br_series, horizons=key_h)
        row    = {'date': dt}
        for h in ws.horizons:
            h_lbl = f'+{h}{u}'
            row[f'edge_{h}{u}']   = float(result.loc[h_lbl, 'edge']) if h_lbl in result.index else np.nan
            row[f'prob_{h}{u}']   = float(result.loc[h_lbl, 'combined']) if h_lbl in result.index else np.nan
            ev = events[h]
            row[f'actual_{h}{u}'] = float(ev.loc[dt]) if dt in ev.index and pd.notna(ev.loc[dt]) else np.nan
        results.append(row)

    df = pd.DataFrame(results)

    print(f"\n=== Backtest [{ws.dir.name}] — {ws.sample_freq} sampling, N={len(df)}, {len(selected)} signals ===")
    print(f"  Nodes: {', '.join(n['id'] for n in selected)}\n")

    pvt   = ws.pivot_horizon
    pvt_n = _h_num(pvt)
    disp_h = ws.display_horizons

    # Calibration, not accuracy. The combiner emits a probability that a barrier is
    # touched; the honest question is whether that number is *right* — when it says
    # 40%, does the barrier get touched 40% of the time? "Directional accuracy"
    # (was the sign correct) is meaningless for a touch event, and against a base
    # rate far from 50% it also flatters a model that only ever predicts the
    # majority class.
    print(f"  Calibration — predicted P(touch) vs realized frequency:\n")
    edges  = [0, 10, 20, 30, 40, 50, 60, 70, 80, 90, 100]
    h_hdr  = ''.join(f'{h:>18}' for h in disp_h)
    print(f"  {'predicted':<14}  {h_hdr}")
    print(f"  {'-'*14}  {''.join('-'*18 for _ in disp_h)}")

    briers = {}
    for lo, hi in zip(edges[:-1], edges[1:]):
        cells, any_row = [], False
        for h in disp_h:
            h_n = _h_num(h)
            sub = df.dropna(subset=[f'prob_{h_n}{u}', f'actual_{h_n}{u}'])
            sub = sub[(sub[f'prob_{h_n}{u}'] >= lo) & (sub[f'prob_{h_n}{u}'] < hi)]
            if len(sub) < 5:
                cells.append(f"{'—':>18}")
                continue
            any_row = True
            realized = sub[f'actual_{h_n}{u}'].mean() * 100
            cells.append(f"{realized:>10.0f}% (n={len(sub):>3})")
        if any_row:
            print(f"  {f'{lo}-{hi}%':<14}  {''.join(cells)}")

    # Brier score against the base rate: does conditioning beat just quoting p0?
    print(f"\n  Brier score (lower is better) — model vs always-quote-the-base-rate:\n")
    print(f"  {'horizon':<8}  {'base rate':>10}  {'model':>8}  {'base':>8}  {'skill':>8}  {'N':>5}")
    print(f"  {'-'*8}  {'-'*10}  {'-'*8}  {'-'*8}  {'-'*8}  {'-'*5}")
    for h in disp_h:
        h_n = _h_num(h)
        sub = df.dropna(subset=[f'prob_{h_n}{u}', f'actual_{h_n}{u}'])
        if sub.empty:
            continue
        br_v   = float(br_series.loc[h]) if h in br_series.index else np.nan
        actual = sub[f'actual_{h_n}{u}']
        model  = ((sub[f'prob_{h_n}{u}'] / 100.0 - actual) ** 2).mean()
        basel  = ((br_v / 100.0 - actual) ** 2).mean()
        # Brier skill score: >0 means conditioning on the features beat the base rate.
        skill  = (1 - model / basel) if basel > 0 else np.nan
        briers[h] = skill
        print(f"  {h:<8}  {br_v:>9.1f}%  {model:>8.4f}  {basel:>8.4f}  {skill:>+8.3f}  {len(sub):>5}")

    if briers and all(v <= 0 for v in briers.values() if pd.notna(v)):
        print("\n  * No horizon beats the base rate. The combined signal adds nothing here.")

    # today's estimate
    t_contribs: list = []
    t_weights:  dict = {}
    for node in selected:
        valid = feat_map[node['id']].dropna()
        if valid.empty:
            continue
        devs   = cmb.eval_at(fn_map[node['id']], float(valid.iloc[-1]))
        n_val  = int(devs['n']) if pd.notna(devs.get('n', np.nan)) else 0
        devs_h = devs[[c for c in devs.index if c.startswith('+')]]
        t_contribs.append((node['id'], devs_h))
        t_weights[node['id']] = cmb.shrink_weight(n_val)

    if t_contribs:
        today = cmb.combine(t_contribs, br_series, horizons=key_h, weights=t_weights)
        print(f"\n  Today's estimate — P({ws.outcome_expr}):\n")
        print(f"  {'horizon':<8}  {'base rate':>10}  {'combined':>10}  {'edge':>10}")
        print(f"  {'-'*8}  {'-'*10}  {'-'*10}  {'-'*10}")
        for h in disp_h:
            if h not in today.index:
                continue
            print(f"  {h:<8}  {float(today.loc[h, 'base_rate']):>9.1f}%  "
                  f"{float(today.loc[h, 'combined']):>9.1f}%  {float(today.loc[h, 'edge']):>+10.1f}pp")
    print()


def cmd_regen(ws: Workspace, family_name: str | None, node_id: str | None) -> None:
    tree = load_tree(ws.tree_path)

    if node_id:
        nodes = [find_node(tree, node_id)]
    elif family_name:
        nodes = all_in_family(tree, family_name)
        if not nodes:
            print(f"No nodes found in family '{family_name}'.")
            return
    else:
        nodes = all_nodes(tree)

    print(f"Regenerating {len(nodes)} node(s)...")
    failed = []
    for node in nodes:
        try:
            run_node(ws, node['id'], regen=True)
        except Exception as e:
            print(f"  [error] {node['id']}: {e} — will retry")
            failed.append(node['id'])

    if failed:
        print(f"\nRetrying {len(failed)} failed node(s)...")
        for nid in failed:
            run_node(ws, nid, regen=True)


# Ordered weakest to strongest; a family inherits the best verdict its node earned
# at any reported horizon.
_VERDICT_RANK = ['insufficient', 'noise', 'underpowered', 'nominal', 'structure']


def _load_validation(ws: Workspace) -> dict:
    """
    Read the validation file matching the active outcome, if one exists.

    Returns {node_id: {horizon: record}}; empty when the sweep has not been run,
    which is itself meaningful — findings then report 'unvalidated' rather than
    claiming an edge no null test has seen.
    """
    path = ws.validation_path
    if not path.exists():
        return {}
    try:
        with open(path, encoding='utf-8') as f:
            raw = json.load(f)
    except (json.JSONDecodeError, OSError):
        return {}
    if raw.get('outcome', {}).get('name', 'up') != ws.outcome:
        return {}
    return {nid: rec.get('horizons', {}) for nid, rec in raw.get('nodes', {}).items()}


def _family_verdict(node_validation: dict) -> tuple[str, dict]:
    """
    Collapse one node's per-horizon null tests into a family verdict.

    The family takes the strongest verdict any horizon earned, together with the
    smallest p-value seen, so a node significant at a single horizon is neither
    hidden nor promoted to significance everywhere.
    """
    if not node_validation:
        return 'unvalidated', {}
    best_rank = -1
    best      = 'noise'
    min_p     = None
    detail    = {}
    for h, rec in node_validation.items():
        v = rec.get('verdict', 'noise')
        detail[h] = {
            'real':     rec.get('real'),
            'null_p95': rec.get('null_p95'),
            'p_value':  rec.get('p_value'),
            'verdict':  v,
        }
        rank = _VERDICT_RANK.index(v) if v in _VERDICT_RANK else 0
        if rank > best_rank:
            best_rank, best = rank, v
        pv = rec.get('p_value')
        if pv is not None and (min_p is None or pv < min_p):
            min_p = pv
    return best, {'min_p_value': min_p, 'horizons': detail}


def cmd_findings(ws: Workspace) -> None:
    by_family = _tested_by_family(ws)
    best      = _select_best_nodes(ws, by_family)

    # Load existing findings to preserve hand-written notes. Like validation, the
    # file is per-outcome so a drawdown sweep never overwrites the directional one.
    findings_path = ws.findings_path
    existing: dict = {}
    if findings_path.exists():
        with open(findings_path, encoding='utf-8') as f:
            raw = json.load(f)
        existing = raw.get('families', {})

    validation = _load_validation(ws)

    read_h   = ws.display_horizons
    families = {}

    for fam in sorted(best):
        best_node, best_pk = best[fam]

        path = ws.output_path(fam, best_node['id'])
        wb   = openpyxl.load_workbook(path, data_only=True)
        conditions = []

        for kind, cond, n_raw, devs in _significant_rows(wb, read_h, ws.read_min_n, ws.read_min_dev):
            conditions.append({
                'direction':  kind,
                'condition':  str(cond),
                'n':          int(n_raw),
                'deviations': {h: round(float(v), 2) for h, v in devs.items()},
            })

        verdict, val_detail = _family_verdict(validation.get(best_node['id'], {}))

        families[fam] = {
            'best_node':   best_node['id'],
            'peak_signal': round(best_pk, 2),
            # 'verdict' is the null-test result, not a count of significant rows:
            # a peak deviation only means something relative to the null that the
            # same max-over-bins statistic produces by chance.
            'verdict':     verdict,
            'has_rows':    bool(conditions),
            'validation':  val_detail,
            'conditions':  conditions,
            'notes':       existing.get(fam, {}).get('notes', ''),
        }

    out = {
        'workspace':        ws.dir.name,
        'generated':        datetime.now(timezone.utc).isoformat(),
        'outcome':          {'name': ws.outcome, 'params': ws.outcome_params,
                             'expr': ws.outcome_expr},
        'display_horizons': read_h,
        'validated':        bool(validation),
        'families':         families,
    }

    with open(findings_path, 'w', encoding='utf-8') as f:
        json.dump(out, f, indent=2)

    print(f"Wrote findings for {len(families)} families → {findings_path.relative_to(ROOT)}")
    if not validation:
        print('  No validation file for this outcome — every family reported as '
              "'unvalidated'. Run --validate first.")
    for v in reversed(_VERDICT_RANK + ['unvalidated']):
        fams = [f for f, rec in families.items() if rec['verdict'] == v]
        if fams:
            print(f"  {v:<13}: {', '.join(sorted(fams))}")


# ── entry ─────────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='Barrier-touch probability pipeline: P(price touches ±theta within h bars | condition)')
    parser.add_argument('--workspace', metavar='NAME', default='btc_daily_14days',
                        help='Workspace name (default: btc_daily_14days)')
    parser.add_argument('--regen',    action='store_true', help='Regenerate xlsx for already-tested nodes')
    parser.add_argument('--rerun',    action='store_true', help='With --family, run every node, not just pending ones')
    parser.add_argument('--family',   metavar='NAME',      help='Target family')
    parser.add_argument('--node',     metavar='ID',        help='Target node')
    parser.add_argument('--families', metavar='FAM,...',   help='Comma-separated families for --probe/--backtest/--validate/--skew')
    parser.add_argument('--outcome',  metavar='NAME',
                        help='Barrier to measure instead of the workspace default: '
                             + ', '.join(outcomes.available()))
    parser.add_argument('--threshold', type=float, metavar='FRAC',
                        help='Barrier size as a fraction of price, e.g. 0.10. Scale it to the '
                             'asset and horizon: 10%% over 14 BTC days is a 21.7%% event, over '
                             '24 NASDAQ hours a 0.4%% one.')
    parser.add_argument('--shifts',   type=int, default=val.DEFAULT_SHIFTS, metavar='N',
                        help='Circular shifts per --validate test (default %d)' % val.DEFAULT_SHIFTS)

    group = parser.add_mutually_exclusive_group()
    group.add_argument('--next',     action='store_true', help='Run the next pending node')
    group.add_argument('--list',     action='store_true', help='List pending nodes by family')
    group.add_argument('--status',   action='store_true', help='Show status table by family')
    group.add_argument('--read',     metavar='ID',        help='Print matrix summary for a node')
    group.add_argument('--probe',    action='store_true', help='Combined touch-probability probe across all families')
    group.add_argument('--skew',     action='store_true',
                       help='Compare both barriers on the same bars: does a condition tilt '
                            'downside specifically, or just raise both?')
    group.add_argument('--backtest', action='store_true', help='Calibration of the combined probability on historical sample dates')
    group.add_argument('--findings', action='store_true', help='Generate findings.<event>.json for the workspace')
    group.add_argument('--validate', action='store_true', help='Shuffle-null test every tested node; writes validation.<event>.json')
    group.add_argument('--charts', action='store_true',
                       help='Regenerate README figures into assets/ from committed data (ignores --workspace)')

    args = parser.parse_args()

    if args.charts:
        print('Rendering README figures into assets/ ...')
        charts.render_all()
        sys.exit(0)

    ws = Workspace(args.workspace)

    if args.outcome or args.threshold is not None:
        name   = args.outcome or ws.outcome
        params = dict(ws.outcome_params) if name == ws.outcome else {}
        if args.threshold is not None:
            params['threshold'] = args.threshold
        try:
            ws.set_outcome(name, params)
        except KeyError as e:
            parser.error(str(e))

    fam_filter = [f.strip() for f in args.families.split(',')] if args.families else None

    if args.read:
        cmd_read(ws, args.read)
    elif args.probe:
        cmd_probe(ws, fam_filter)
    elif args.skew:
        cmd_skew(ws, fam_filter)
    elif args.backtest:
        cmd_backtest(ws, fam_filter)
    elif args.findings:
        cmd_findings(ws)
    elif args.validate:
        cmd_validate(ws, fam_filter, n_shifts=args.shifts)
    elif args.regen:
        cmd_regen(ws, args.family, args.node)
    elif args.next:
        cmd_next(ws)
    elif args.list:
        cmd_list(ws)
    elif args.status:
        cmd_status(ws)
    elif args.family:
        cmd_family(ws, args.family, rerun=args.rerun)
    elif args.node:
        run_node(ws, args.node)
    else:
        parser.print_help()
