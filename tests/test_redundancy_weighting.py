from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from barrierlab.domain import bayes, redundancy
from barrierlab.infrastructure import artifact_io
from barrierlab.pipeline.step_05_redundancy import validation_fingerprint
from barrierlab.presentation import workbooks


class RedundancyTests(unittest.TestCase):
    def test_conditional_nmi_identifies_duplicate_bin_states(self) -> None:
        left = np.tile(np.arange(10), 40)
        y = np.tile(np.array([0, 1]), 200)
        self.assertAlmostEqual(redundancy.conditional_nmi(left, left, y), 1.0)

    def test_clusters_keep_partial_information_nodes_without_hard_deletion(self) -> None:
        similarity = np.array([
            [1.00, 0.45, 0.05],
            [0.45, 1.00, 0.10],
            [0.05, 0.10, 1.00],
        ])
        self.assertEqual(redundancy.clusters(similarity, 0.30), [[0, 1], [2]])

    def test_numerical_artifact_round_trips_without_pickle(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "redundancy.npz"
            artifact_io.save_redundancy(path, {
                "node_ids": np.array(["atr", "rv"], dtype=str),
                "conditional_nmi": np.array([[1.0, 0.8], [0.8, 1.0]]),
                "cluster_id": np.array([1, 1]),
            }, {"source_validation_fingerprint": "abc"})

            result = artifact_io.load_redundancy(path)

        np.testing.assert_array_equal(result["node_ids"], np.array(["atr", "rv"]))
        np.testing.assert_allclose(result["conditional_nmi"], [[1.0, 0.8], [0.8, 1.0]])
        self.assertEqual(result["meta"]["source_validation_fingerprint"], "abc")

    def test_workbook_exposes_matrix_clusters_nodes_and_pairs(self) -> None:
        clusters = [{
            "cluster": 1,
            "representative": "atr",
            "members": ["atr", "rv"],
            "families": ["volatility"],
            "size": 2,
            "representative_information": 0.03,
            "mean_internal_nmi": 0.8,
            "max_internal_nmi": 0.8,
        }]
        pairs = [{
            "left": "atr", "right": "rv", "conditional_nmi": 0.8,
            "above_threshold": True, "same_cluster": True,
        }]
        validation = {
            node: {
                "selection_rank": rank, "horizon": 14, "peak_p": 0.001,
                "q_value": 0.002, "economic": {"best": {"dev": 15.0, "Δ": -0.1}},
            }
            for rank, node in enumerate(["atr", "rv"], 1)
        }
        meta = {
            "workspace": "example",
            "generated": "2026-01-01T00:00:00+00:00",
            "target": {"Δ": -0.1, "horizon": 14, "unit": "d"},
            "method": {"threshold": 0.3},
            "n_observations": 100,
            "outcome_rate": 0.25,
            "source_validation_fingerprint": "abc",
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "redundancy.xlsx"
            workbooks.write_redundancy_xlsx(
                path,
                ["atr", "rv"],
                ["volatility", "volatility"],
                np.array([0.03, 0.02]),
                np.array([[1.0, 0.8], [0.8, 1.0]]),
                np.array([1, 1]),
                np.array([True, False]),
                clusters,
                pairs,
                validation,
                meta,
            )
            workbook = load_workbook(path, read_only=True)
            sheets = workbook.sheetnames
            workbook.close()

        self.assertEqual(
            sheets,
            ["Overview", "Matrix", "Clusters", "Nodes", "Pairs", "Definitions"],
        )

    def test_validation_fingerprint_tracks_cleared_decisions(self) -> None:
        validation = {
            "selection_fingerprint": [{"node": "atr"}],
            "method": {"q": 0.05},
            "cleared": [{"node": "atr", "q_value": 0.01}],
        }
        changed = {
            **validation,
            "cleared": [{"node": "atr", "q_value": 0.02}],
        }

        self.assertEqual(validation_fingerprint(validation), validation_fingerprint(validation))
        self.assertNotEqual(validation_fingerprint(validation), validation_fingerprint(changed))


class WeightedBayesTests(unittest.TestCase):
    def test_ridge_keeps_unique_information_in_overlapping_nodes(self) -> None:
        rng = np.random.default_rng(17)
        shared = rng.normal(size=5000)
        unique = rng.normal(size=5000)
        atr = shared
        realized_vol = 0.8 * shared + 0.6 * unique
        logits = -0.3 + 0.9 * atr + 0.4 * realized_vol
        y = rng.binomial(1, 1.0 / (1.0 + np.exp(-logits)))
        model = bayes.fit_weighted_logistic(
            np.column_stack([atr, realized_vol]), y, ridge=0.10
        )

        self.assertGreater(model["weights"][0], 0.05)
        self.assertGreater(model["weights"][1], 0.05)
        self.assertTrue(np.all(model["weights"] >= 0))

    def test_stage_six_workbook_is_one_stage_one_formatted_probability_sheet(self) -> None:
        probability = np.array([[0.20, 0.40], [0.80, 0.90]])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bayes.xlsx"
            workbooks.write_bayes_xlsx(
                path,
                probability,
                np.array([-0.1, 0.1]),
                np.array([1, 7]),
                np.array([500, 494]),
            )
            workbook = load_workbook(path, read_only=False)
            sheet = workbook["Weighted Bayes"]
            self.assertEqual(workbook.sheetnames, ["Weighted Bayes"])
            self.assertEqual(sheet["A1"].value, "Condition —— Weighted Bayes")
            self.assertEqual(sheet["A2"].value, "P (Price touches Δ within t | Condition)")
            self.assertEqual(sheet["A4"].value, "Δ")
            self.assertEqual(sheet["B4"].value, "+1d")
            self.assertEqual(sheet["B3"].value, 500)
            self.assertEqual(sheet["B5"].value, 0.8)
            self.assertEqual(sheet["B5"].number_format, "0.0%")
            self.assertEqual(sheet.freeze_panes, "B5")
            self.assertEqual(len(sheet.conditional_formatting), 1)
            workbook.close()

    def test_stage_six_shift_workbook_is_one_stage_two_formatted_shift_sheet(self) -> None:
        probability = np.array([[0.20, 0.40], [0.80, 0.90]])
        baseline = np.array([[0.25, 0.35], [0.70, 0.95]])
        shift_pp = (probability - baseline) * 100.0
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bayes_shift.xlsx"
            workbooks.write_bayes_shift_xlsx(
                path,
                shift_pp,
                probability,
                baseline,
                np.array([-0.1, 0.1]),
                np.array([1, 7]),
                np.array([500, 494]),
            )
            workbook = load_workbook(path, read_only=False)
            sheet = workbook["Weighted Bayes"]
            self.assertEqual(workbook.sheetnames, ["Weighted Bayes"])
            self.assertEqual(sheet["A1"].value, "Condition —— Weighted Bayes")
            self.assertEqual(
                sheet["A2"].value,
                "P (Price touches Δ within t | Condition) - P (Price touches Δ within t)",
            )
            self.assertEqual(sheet["B5"].value, 0.1)
            self.assertEqual(sheet["B5"].number_format, "+0.0%;-0.0%;0.0%")
            self.assertEqual(sheet.freeze_panes, "B5")
            self.assertEqual(len(sheet.conditional_formatting), 1)
            workbook.close()

    def test_current_surface_obeys_barrier_and_horizon_nesting(self) -> None:
        deltas = np.array([-0.2, -0.1, 0.0, 0.1, 0.2])
        raw = np.array([
            [0.20, 0.10, 0.30],
            [0.15, 0.40, 0.35],
            [0.80, 0.70, 0.90],
            [0.60, 0.50, 0.80],
            [0.70, 0.30, 0.75],
        ])
        result = bayes.coherent_probability_surface(raw, deltas)

        self.assertTrue(np.all(np.diff(result, axis=1) >= -1e-9))
        self.assertTrue(np.all(np.diff(result[:3], axis=0) >= -1e-9))
        self.assertTrue(np.all(np.diff(result[2:], axis=0) <= 1e-9))
        self.assertTrue(np.all((result >= 0.0) & (result <= 1.0)))

    def test_current_forecast_uses_latest_state_but_only_completed_labels(self) -> None:
        rng = np.random.default_rng(23)
        index = pd.date_range("2024-01-01", periods=500, freq="D")
        close = 100.0 * np.exp(np.cumsum(rng.normal(0.0, 0.01, len(index))))
        data = pd.DataFrame({
            "close": close,
            "high": close * (1.0 + rng.uniform(0.002, 0.02, len(index))),
            "low": close * (1.0 - rng.uniform(0.002, 0.02, len(index))),
        }, index=index)
        feats = {
            "momentum": pd.Series(rng.normal(size=len(index)), index=index),
            "volatility": pd.Series(rng.lognormal(size=len(index)), index=index),
        }

        result = bayes.current_weighted_forecast(
            data, feats, -0.02, 7, n_bins=10, top_k=2
        )
        historical = bayes.current_weighted_forecast(
            data, feats, -0.02, 7, n_bins=10, top_k=2, as_of="2025-03-01"
        )

        self.assertEqual(result["as_of"], "2025-05-14")
        self.assertLessEqual(result["n_observations"], len(index) - 7)
        self.assertGreaterEqual(result["probability"], 0.0)
        self.assertLessEqual(result["probability"], 1.0)
        self.assertEqual(set(result["selected_nodes"]), set(feats))
        self.assertEqual(historical["as_of"], "2025-03-01")
        self.assertLess(historical["n_observations"], result["n_observations"])

        surface = bayes.current_weighted_surface(
            data,
            feats,
            np.array([-0.02, 0.0, 0.02]),
            np.array([1, 7]),
            n_bins=10,
            top_k=2,
            as_of="2025-03-01",
        )
        self.assertEqual(surface["as_of"], "2025-03-01")
        self.assertEqual(surface["probability"].shape, (3, 2))
        self.assertEqual(surface["raw_probability"].shape, (3, 2))
        self.assertEqual(surface["n_observations"].shape, (2,))
        self.assertTrue(np.isfinite(surface["probability"]).all())


if __name__ == "__main__":
    unittest.main()
